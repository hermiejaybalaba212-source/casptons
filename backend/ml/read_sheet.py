#!/usr/bin/env python
"""Read a CSV / Excel upload and print every row and column as JSON.

    python ml/read_sheet.py <file>

Output (one JSON object):
    {"ok": true, "columns": [...], "rows": [{col: value, ...}, ...],
     "row_count": n, "sheet": "Sheet1", "years": ["AY 2025-2026", ...]}

Nothing is assumed about the file: the header is *found* (title rows above it
are skipped), and any school-year label that appears anywhere in the sheet is
reported so the UI can group by whatever years the file actually contains.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional

# AY 2025-2026 / SY 2025/2026 / 2025-2026 / A.Y. 2025 - 2026, optionally with
# a trailing school-year tag glued on ("AY2026-2027 (2nd Sem)"). Boundaries
# exclude alphanumerics so glued filenames still match - see YEAR_EXPLICIT_RE.
YEAR_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:A\.?\s?Y\.?|S\.?\s?Y\.?|AY|SY|academic\s+year|school\s+year)?\s*"
    r"(20\d{2})\s*[-–/]\s*(20\d{2})(?!\d)",
    re.I,
)

# Same range, but only when the cell *says* it is a school year. An ID column
# holding "2013-2014" or a student_id of "2024-2025" must never become a
# category - only a real AY/SY label does.
# Both boundaries exclude alphanumerics rather than using \b: "_" is a word
# character, so \b would reject the glued labels real files carry -
# "spc_approved_AY2026-2027.csv" and "approved_AY2024-2025_final.xlsx" - while
# still needing to reject "friday 2025-2026" (the "ay" in Friday) and
# "20251-20262" (a stray longer number).
YEAR_EXPLICIT_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:A\.?\s?Y\.?|S\.?\s?Y\.?|AY|SY|academic\s+year|school\s+year)\s*[-–/]?\s*"
    r"(20\d{2})\s*[-–/]\s*(20\d{2})(?!\d)",
    re.I,
)

# Column names that are allowed to carry a bare "2025-2026" without an AY/SY
# prefix (source_file, remarks, batch, term, semester, ...). Identity columns
# are deliberately absent: student_id / lrn / philsys_id hold ID fragments.
YEAR_COLUMN_HINTS = (
    "year", "ay_", "_ay", "sy_", "_sy", "academic", "school", "term",
    "semester", "sem", "source", "remark", "batch", "period", "aye",
)


def _column_may_hold_year(column: str) -> bool:
    name = column.strip().lower()
    return any(hint in name for hint in YEAR_COLUMN_HINTS)


ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
MAX_PREVIEW_ROWS = 5000


def out(payload: dict) -> None:
    sys.stdout.write(json.dumps(payload, ensure_ascii=False, default=str))
    sys.stdout.flush()


def canonical_year(value: str, column: str = "") -> str:
    """Normalise a school-year cell to the AY 2024-2025 label the UI shows.

    Accepts an explicit AY/SY label anywhere, or a bare range when the column
    name itself is year-like. Anything else - an ID, a birthdate - returns ""
    so it cannot be mistaken for a school year.
    """
    text = str(value)
    match = YEAR_EXPLICIT_RE.search(text)
    if match is None and column and _column_may_hold_year(column):
        match = YEAR_RE.search(text)
    if match is None:
        return ""
    return f"AY {match.group(1)}-{match.group(2)}"


def scan_years(rows: List[Dict[str, str]], columns: List[str]) -> tuple[List[str], str]:
    """Every school year present anywhere in the sheet, newest first.

    Also returns the column that actually carried them, so the UI can group by
    it instead of re-scanning every cell.
    """
    found: Dict[str, set] = {col: set() for col in columns}
    for row in rows:
        for col in columns:
            year = canonical_year(row.get(col, ""), col)
            if year:
                found[col].add(year)

    best_col, best = "", set()
    for col in columns:
        if len(found[col]) > len(best):
            best_col, best = col, found[col]
    years = sorted(set().union(*found.values()) if found else set(), reverse=True)
    return years, best_col


# A "year_level" / "level" column carries the student's current level (1, 2, 3,
# 4) rather than a school year. Applicants lists shipped by the registrar use it
# instead of an AY column, so it must still section the table by year.
YEAR_LEVEL_COLUMN_RE = re.compile(r"^(?:year|yr)[\s_-]*level$|^level$", re.I)
YEAR_LEVEL_VALUE_RE = re.compile(r"^\s*(\d{1,2})\s*(?:st|nd|rd|th)?\s*(?:year|level)?\s*$", re.I)


def find_year_level_column(columns: List[str]) -> str:
    """The year-level column ("year_level", "Year Level", "level"), if any."""
    for col in columns:
        if YEAR_LEVEL_COLUMN_RE.match(col.strip()):
            return col
    return ""


def ordinal_year(n: int) -> str:
    """3 -> "3rd Year" (the label the UI shows and groups by)."""
    mod100 = n % 100
    if 11 <= mod100 <= 13:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix} Year"


def scan_year_levels(rows: List[Dict[str, str]], column: str) -> Dict[str, int]:
    """Distinct year levels in the column -> how many rows carry each.

    Values that are not a plain level ("", "N/A", a name) are ignored so a
    messy column cannot invent a bogus category.
    """
    counts: Dict[str, int] = {}
    for row in rows:
        match = YEAR_LEVEL_VALUE_RE.match(text(row.get(column, "")))
        if not match:
            continue
        label = ordinal_year(int(match.group(1)))
        counts[label] = counts.get(label, 0) + 1
    return counts


def guess_group_column(columns: List[str], year_column: str) -> str:
    """Column to section the table by when the file carries no school year.

    A program/course column is the natural category for an applicants list;
    without one there is nothing sensible to split on, so return "" and let the
    UI show a single flat table.
    """
    if year_column:
        return year_column
    hints = ("program", "course", "degree", "scholarship", "grant", "award")
    lowered = [(col, col.strip().lower()) for col in columns]
    for hint in hints:
        for col, name in lowered:
            if hint in name:
                return col
    return ""


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def pick_header(rows: List[List[str]]) -> tuple[List[str], int]:
    """Title rows sit above the real header on every Annex/billing sheet.

    The header row is the one with the most populated cells - a title row has
    1-2 words, a header has 5+. Fall back to the first non-empty row.
    """
    best_index, best_score = -1, 0
    for index, row in enumerate(rows[:25]):
        cells = [c for c in row if c]
        if not cells:
            continue
        # Prefer many populated cells; break ties by how label-like they are.
        labelish = sum(1 for c in cells if not c.replace(",", "").replace(".", "").isdigit())
        score = len(cells) * 10 + labelish
        if score > best_score:
            best_index, best_score = index, score
    if best_index < 0:
        return [], -1

    header = rows[best_index]
    seen: Dict[str, int] = {}
    columns: List[str] = []
    for position, cell in enumerate(header):
        name = cell or f"Column {position + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        columns.append(name)
    return columns, best_index


def read_csv(path: Path) -> tuple[List[str], List[Dict[str, str]]]:
    raw = path.read_bytes()
    text_content = ""
    for encoding in ENCODINGS:
        try:
            text_content = raw.decode(encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if not text_content:
        return [], []

    sample = text_content[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    grid = [row for row in csv.reader(io.StringIO(text_content), dialect)]
    return _grid_to_records(grid)


def read_xlsx(path: Path) -> tuple[List[str], List[Dict[str, str]], str]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("openpyxl is not installed - save the sheet as CSV instead")

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        # Pick the sheet with the most populated rows - Annex workbooks carry a
        # blank cover/legend sheet before the real data.
        chosen, chosen_grid, chosen_score = None, [], -1
        for sheet in book.worksheets:
            grid = [[text(c) for c in row] for row in sheet.iter_rows(values_only=True)]
            score = sum(1 for row in grid if any(row))
            if score > chosen_score:
                chosen, chosen_grid, chosen_score = sheet.title, grid, score
        if chosen is None:
            return [], [], ""
        columns, header_index = pick_header(chosen_grid)
        rows = _to_records(chosen_grid, columns, header_index)
        return columns, rows, chosen
    finally:
        book.close()


def _grid_to_records(grid: List[List[str]]) -> tuple[List[str], List[Dict[str, str]]]:
    columns, header_index = pick_header(grid)
    if header_index < 0:
        return [], []
    return columns, _to_records(grid, columns, header_index)


def _to_records(grid: List[List[str]], columns: List[str], header_index: int) -> List[Dict[str, str]]:
    records: List[Dict[str, str]] = []
    for row in grid[header_index + 1:]:
        if not any(row):
            continue
        record: Dict[str, str] = {}
        for position, column in enumerate(columns):
            record[column] = text(row[position]) if position < len(row) else ""
        # Drop trailing blank rows that spreadsheets leave behind.
        if any(record.values()):
            records.append(record)
        if len(records) >= MAX_PREVIEW_ROWS:
            break
    return records


def main(argv: List[str]) -> int:
    if len(argv) < 2:
        out({"ok": False, "error": "usage: read_sheet.py <file>"})
        return 1

    path = Path(argv[1])
    if not path.is_file():
        out({"ok": False, "error": f"file not found: {path}"})
        return 1

    suffix = path.suffix.lower()
    try:
        if suffix == ".csv":
            columns, rows = read_csv(path)
            sheet = path.name
        elif suffix in (".xlsx", ".xlsm", ".xltx"):
            columns, rows, sheet = read_xlsx(path)
        elif suffix == ".xls":
            out({"ok": False, "error": "Legacy .xls is not supported - re-save as .xlsx or .csv"})
            return 1
        else:
            out({"ok": False, "error": f"Unsupported file type '{suffix}' - use .csv or .xlsx"})
            return 1
    except Exception as exc:  # surfaced as {error} to the UI
        out({"ok": False, "error": str(exc)})
        return 1

    years, year_column = scan_years(rows, columns)

    # A real AY/SY column wins. When the file carries none, fall back to the
    # year-level column ("1", "2", "3", "4") so an applicants list without a
    # school-year column still sections by year instead of by program.
    year_kind = "ay" if years else ""
    year_counts: Dict[str, int] = {}
    if years:
        # Count rows per canonical AY label from the column that carried them.
        for row in rows:
            label = canonical_year(row.get(year_column, ""), year_column)
            if label:
                year_counts[label] = year_counts.get(label, 0) + 1
    else:
        level_col = find_year_level_column(columns)
        if level_col:
            level_counts = scan_year_levels(rows, level_col)
            if level_counts:
                # Ascending: 1st Year first, the order a class list reads in.
                def _level_number(label: str) -> int:
                    match = re.match(r"^(\d+)", label)
                    return int(match.group(1)) if match else 0

                years = sorted(level_counts, key=_level_number)
                year_column = level_col
                year_kind = "level"
                year_counts = level_counts

    out(
        {
            "ok": True,
            "columns": columns,
            "rows": rows,
            "row_count": len(rows),
            "sheet": sheet,
            "years": years,
            "year_column": year_column,
            # "ay" = school-year labels (AY 2025-2026), "level" = year level
            # (1st Year, 2nd Year). The UI labels each section accordingly.
            "year_kind": year_kind,
            "year_counts": year_counts,
            # The column the UI should section on: the year column when the
            # file carries one, otherwise the program/course column.
            "group_column": guess_group_column(columns, year_column),
        }
    )
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8") if hasattr(sys.stdout, "reconfigure") else None
    sys.exit(main(sys.argv))
